from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

OUTPUT = Path(__file__).resolve().parents[1] / "public" / "Oakridge-MUN-Registration-Test.xlsx"

HEADERS = [
    "Response ID",
    "Start time",
    "Completion time",
    "Delegate Full Name",
    "Student Email ID",
    "Student Phone Number",
    "Name of School",
    "Grade",
    "Previous MUN Experience",
    "Committee Preference 1",
    "Committee Preference 2",
    "Committee Preference 3",
    "Country Preference 1",
    "Country Preference 2",
    "Country Preference 3",
    "Payment Status",
    "Receipt Number",
    "Do you have any questions for us?",
]

COMMITTEES = [
    "DISEC",
    "ARMAGEDDON",
    "JCC-1",
    "JCC-2",
    "UNSC",
    "UNHRC",
    "OIC",
    "UNEP",
    "WHO",
    "Lok Sabha",
]

SCHOOLS = [
    "Oakridge International School",
    "CHIREC International School",
    "Indus International School",
    "Delhi Public School Hyderabad",
]

PEOPLE = [
    ("Aanya Verma", "aanya.verma@example.com", "DISEC", "UNHRC", "UNSC", "France", "Brazil", "Japan"),
    ("Arjun Rao", "arjun.rao@example.com", "ARMAGEDDON", "JCC-1", "UNSC", "Research Lead", "Systems Architect", "Security Director"),
    ("Meera Iyer", "meera.iyer@example.com", "DISEC", "UNEP", "OIC", "Kenya", "Germany", "Indonesia"),
    ("Kabir Shah", "kabir.shah@example.com", "JCC-1", "JCC-2", "ARMAGEDDON", "United States", "United Kingdom", "Soviet Union"),
    ("Zoya Khan", "zoya.khan@example.com", "UNHRC", "DISEC", "WHO", "South Africa", "Mexico", "Norway"),
    ("Rohan Nair", "rohan.nair@example.com", "ARMAGEDDON", "UNSC", "JCC-2", "Infrastructure Chief", "Diplomatic Envoy", "Field Commander"),
    ("Tara Menon", "tara.menon@example.com", "DISEC", "OIC", "UNEP", "Colombia", "Nigeria", "Sweden"),
    ("Vihaan Reddy", "vihaan.reddy@example.com", "JCC-2", "JCC-1", "UNSC", "Soviet Union", "France", "China"),
    ("Sara Joseph", "sara.joseph@example.com", "WHO", "UNHRC", "DISEC", "Canada", "India", "Chile"),
    ("Dev Malhotra", "dev.malhotra@example.com", "UNSC", "ARMAGEDDON", "DISEC", "China", "United States", "France"),
    ("Nisha Patel", "nisha.patel@example.com", "OIC", "DISEC", "UNHRC", "UAE", "Türkiye", "Malaysia"),
    ("Reyansh Gupta", "reyansh.gupta@example.com", "ARMAGEDDON", "JCC-2", "DISEC", "Cybersecurity Minister", "Energy Minister", "Public Information Director"),
]

NAVY = "071D49"
BLUE = "0055C9"
CYAN = "21C4D7"
LIGHT = "EAF2FA"
INK = "132238"
WHITE = "FFFFFF"
GREEN = "DDF4E7"
AMBER = "FFF0C2"
RED = "FDE2E2"
THIN = Side(style="thin", color="CDD9E6")


def style_header(ws):
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True, size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=Side(style="medium", color=CYAN))
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "A2"


def style_data_sheet(ws, rows, table_name):
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    style_header(ws)
    widths = [15, 21, 21, 23, 30, 22, 31, 10, 22, 25, 25, 25, 24, 24, 24, 17, 18, 34]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)
    if ws.max_row >= 2:
        table = Table(displayName=table_name, ref=f"A1:R{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
        ws.add_table(table)
        ws.conditional_formatting.add(
            f"I2:I{ws.max_row}",
            ColorScaleRule(start_type="min", start_color="F8FBFF", end_type="max", end_color="7FDBCA"),
        )


def make_response(index, person):
    name, email, pref1, pref2, pref3, country1, country2, country3 = person
    minute = 5 + index * 7
    hour = 9 + minute // 60
    minute %= 60
    start = f"2026-08-09 {hour:02d}:{minute:02d}:00"
    end_minute = minute + 4
    end_hour = hour + end_minute // 60
    end_minute %= 60
    completed = f"2026-08-09 {end_hour:02d}:{end_minute:02d}:00"
    school = SCHOOLS[index % len(SCHOOLS)]
    grade = str(9 + index % 4)
    experience = f"{index % 5} conference" + ("s" if index % 5 != 1 else "")
    payment = ["Paid", "Pending", "Paid", "Unpaid"][index % 4]
    receipt = f"OAK26-{1001 + index}" if payment == "Paid" else ""
    question = "Please confirm accessibility arrangements." if index == 6 else ""
    return [
        f"R-{1042 + index}", start, completed, name, email, f"+91 90000 {10000 + index}", school,
        grade, experience, pref1, pref2, pref3, country1, country2, country3, payment, receipt, question,
    ]


def add_validations(ws, max_row=500):
    committee = DataValidation(type="list", formula1="'Lists'!$A$2:$A$11", allow_blank=True)
    grade = DataValidation(type="list", formula1='"8,9,10,11,12"', allow_blank=True)
    payment = DataValidation(type="list", formula1='"Paid,Pending,Unpaid,Needs review"', allow_blank=True)
    ws.add_data_validation(committee)
    ws.add_data_validation(grade)
    ws.add_data_validation(payment)
    committee.add(f"J2:L{max_row}")
    grade.add(f"H2:H{max_row}")
    payment.add(f"P2:P{max_row}")


def build():
    wb = Workbook()
    responses = wb.active
    assert responses is not None
    responses.title = "Responses"
    rows = [make_response(index, person) for index, person in enumerate(PEOPLE)]
    style_data_sheet(responses, rows, "RegistrationResponses")
    add_validations(responses)

    template = wb.create_sheet("Blank Form Template")
    style_data_sheet(template, [["" for _ in HEADERS]], "BlankRegistrationTemplate")
    add_validations(template)

    edge = wb.create_sheet("Edge Case Tests")
    duplicate = make_response(0, PEOPLE[0])
    duplicate[0] = "R-EDGE-DUPLICATE"
    duplicate[3] = "Aanya V. Duplicate"
    missing_email = make_response(1, PEOPLE[1])
    missing_email[0] = "R-EDGE-NO-EMAIL"
    missing_email[4] = ""
    missing_preference = make_response(2, PEOPLE[2])
    missing_preference[0] = "R-EDGE-NO-PREFERENCE"
    missing_preference[9:12] = ["", "", ""]
    style_data_sheet(edge, [duplicate, missing_email, missing_preference], "RegistrationEdgeCases")
    add_validations(edge)

    readme = wb.create_sheet("Read Me")
    readme.sheet_view.showGridLines = False
    readme.column_dimensions["A"].width = 4
    readme.column_dimensions["B"].width = 34
    readme.column_dimensions["C"].width = 92
    readme.merge_cells("B2:C2")
    readme["B2"] = "OAKRIDGE MUN — REGISTRATION SYSTEM TEST KIT"
    readme["B2"].font = Font(size=20, bold=True, color=WHITE)
    readme["B2"].fill = PatternFill("solid", fgColor=NAVY)
    readme["B2"].alignment = Alignment(vertical="center")
    readme.row_dimensions[2].height = 48
    instructions = [
        ("1 · Clean import", "Import the Responses sheet through Operations → Forms. It contains 12 fictional delegates and example.com addresses. No real person will be contacted."),
        ("2 · Contacts check", "After import, open Contacts. Every valid Student Email ID should appear once with preferences and the source file recorded."),
        ("3 · Demand check", "The Forms workspace should combine first, second, and third choices into one demand view—there are no round tabs."),
        ("4 · Allocation check", "Enter committee capacities, generate recommendations, review them, then explicitly apply accepted recommendations to Contacts."),
        ("5 · Parser stress test", "To test errors, copy the Edge Case Tests sheet into a new workbook as the first sheet. It includes a duplicate email, missing email, and missing preferences."),
        ("Privacy", "All names, phone numbers, emails, response IDs, and receipt numbers are fictional test data. The workbook is safe to keep in the repository."),
    ]
    for row_index, (title, body) in enumerate(instructions, 4):
        readme[f"B{row_index}"] = title
        readme[f"B{row_index}"].font = Font(bold=True, color=NAVY)
        readme[f"C{row_index}"] = body
        readme[f"C{row_index}"].alignment = Alignment(wrap_text=True, vertical="top")
        readme[f"B{row_index}"].fill = PatternFill("solid", fgColor=LIGHT)
        readme[f"B{row_index}"].border = Border(left=Side(style="medium", color=CYAN))
        readme.row_dimensions[row_index].height = 48
    readme["B12"] = "Use only for testing"
    readme["B12"].font = Font(bold=True, color="8B1E1E")
    readme["C12"] = "Do not send test campaigns to the fictional example.com addresses."
    readme["C12"].fill = PatternFill("solid", fgColor=RED)

    lists = wb.create_sheet("Lists")
    lists.append(["Committees"])
    for committee in COMMITTEES:
        lists.append([committee])
    lists.sheet_state = "hidden"

    wb.properties.title = "Oakridge MUN Registration System Test Kit"
    wb.properties.subject = "Microsoft Forms-compatible registration and contacts test workbook"
    wb.properties.creator = "Oakridge MUN Operations"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)

    check = load_workbook(OUTPUT, read_only=False, data_only=False)
    assert check.sheetnames == ["Responses", "Blank Form Template", "Edge Case Tests", "Read Me", "Lists"]
    assert check["Responses"].max_row == len(PEOPLE) + 1
    assert check["Responses"]["A2"].value == "R-1042"
    assert check["Responses"]["J2"].value == "DISEC"
    print(f"created={OUTPUT}")
    print(f"responses={len(PEOPLE)} sheets={len(check.sheetnames)}")


if __name__ == "__main__":
    build()
